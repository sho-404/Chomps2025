"""Fill the IEEE Access "Response to Reviewers" template from a JSON file.

Usage:
    python code/common/make_response_docx.py --template <template.docx> --content <content.json> --out <out.docx>

content.json:
{
  "manuscript_id": "Access-2026-38324",
  "title": "...",
  "authors": "Vega Guzmán et al.",
  "summary": ["paragraph", ...],              # optional paragraphs after the greeting
  "reviewers": [
    {"name": "Reviewer#1", "concerns": [
        {"title": "...", "comment": "...", "response": "...", "action": "..."}, ...]},
    ...
  ]
}
The template's own example concern blocks are removed; its header, greeting,
closing and trailing note are kept.
"""
import argparse
import copy
import json

import docx
from docx.shared import Pt


def _delete_paragraph(p):
    el = p._element
    el.getparent().remove(el)


def _add_par(doc, before_el, text, bold_prefix=None, italic=False):
    """Insert a paragraph before `before_el` (an lxml element)."""
    p = doc.add_paragraph()
    before_el.addprevious(p._element)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(11)
    return p


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--content", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    c = json.load(open(a.content, encoding="utf-8"))
    doc = docx.Document(a.template)
    pars = doc.paragraphs
    # header fields
    for p in pars:
        if p.text.startswith("Original Manuscript ID"):
            for r in p.runs:
                r.text = ""
            p.runs[0].text = f"Original Manuscript ID: {c['manuscript_id']}"
        elif p.text.startswith("Original Article Title"):
            for r in p.runs:
                r.text = ""
            p.runs[0].text = f"Original Article Title: “{c['title']}”"
        elif p.text.strip().startswith("<author name>"):
            for r in p.runs:
                r.text = ""
            p.runs[0].text = f"{c['authors']}"
    # locate the example concern blocks: from the first "Reviewer#" paragraph to the "Note:" paragraph
    first = next(p for p in doc.paragraphs if p.text.startswith("Reviewer#"))
    note = next(p for p in doc.paragraphs if p.text.startswith("Note:"))
    # remove everything between first and note (exclusive of note)
    el = first._element
    while el is not None and el is not note._element:
        nxt = el.getnext()
        el.getparent().remove(el)
        el = nxt
    anchor = note._element
    # optional summary paragraphs right after the greeting block (before the concerns)
    for s in c.get("summary", []):
        _add_par(doc, anchor, s)
    for rev in c["reviewers"]:
        for i, con in enumerate(rev["concerns"], 1):
            _add_par(doc, anchor, "")
            _add_par(doc, anchor, f" {con['title']}", bold_prefix=f"{rev['name']}, Concern # {i}:")
            _add_par(doc, anchor, con["comment"], italic=True)
            _add_par(doc, anchor, " " + con["response"], bold_prefix="Author response:")
            _add_par(doc, anchor, " " + con["action"], bold_prefix="Author action:")
    doc.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
